#!/usr/bin/env python3
"""版面 QA：估算中文文字所需高度，找出溢位、出血、重疊。"""
import sys, math
from pptx import Presentation
from pptx.util import Emu, Length

EMU_IN = 914400.0
SW, SH = 13.333, 7.5

def char_w(ch, size):           # 回傳該字元寬度（英吋）
    o = ord(ch)
    if o > 0x2E80: em = 1.0     # CJK 全形
    elif ch == ' ': em = 0.3
    elif ch.isdigit() or ch.isupper(): em = 0.6
    else: em = 0.52
    return em * size / 72.0

def line_count(text, box_w, size):
    lines = 0
    for seg in text.split("\n"):
        if not seg.strip():
            lines += 1; continue
        w, n = 0.0, 1
        for ch in seg:
            cw = char_w(ch, size)
            if w + cw > box_w:
                n += 1; w = cw
            else:
                w += cw
        lines += n
    return max(lines, 1)

def txt_need(shape):
    tf = shape.text_frame
    box_w = shape.width / EMU_IN - 0.14
    total = 0.0
    paras = tf.paragraphs
    for i, p in enumerate(paras):
        text = "".join(r.text for r in p.runs)
        if not text:
            total += 0.14; continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
        ls = p.line_spacing
        if ls is None:
            lh_pt = 1.2 * size
        elif isinstance(ls, Length):
            lh_pt = ls.pt
        else:
            lh_pt = ls * size
        lh = lh_pt / 72.0
        bw = box_w - (0.22 if p.level is not None and _has_bullet(p) else 0)
        total += line_count(text, max(bw, 0.5), size) * lh
        sa = p.space_after
        if sa is not None and i != len(paras) - 1:
            total += sa.pt / 72.0
    return total

def _has_bullet(p):
    return b'buChar' in p._pPr.xml.encode() if p._pPr is not None else False

def table_need(gf):
    tbl = gf.table
    total = 0.0
    for r in tbl.rows:
        maxl = 1
        for ci, cell in enumerate(r.cells):
            cw = tbl.columns[ci].width / EMU_IN - 0.18
            size = max((run.font.size.pt for para in cell.text_frame.paragraphs
                        for run in para.runs if run.font.size), default=14)
            maxl = max(maxl, line_count(cell.text, max(cw, 0.4), size))
        total += max(r.height / EMU_IN, maxl * size * 1.25 / 72.0 + 0.12)
    return total

prs = Presentation(sys.argv[1])
issues = []
for i, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        L, T = sh.left / EMU_IN, sh.top / EMU_IN
        W, H = sh.width / EMU_IN, sh.height / EMU_IN
        if L < -0.02 or T < -0.02 or L + W > SW + 0.02 or T + H > SH + 0.02:
            issues.append(f"P{i} 出血：{sh.shape_type} @({L:.2f},{T:.2f},{W:.2f}x{H:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            need = txt_need(sh)
            if need > H + 0.03:
                issues.append(f"P{i} 文字溢位：需 {need:.2f}\" / 框 {H:.2f}\" ｜「{sh.text_frame.text.strip()[:26]}…」")
            boxes.append((L, T, W, H, sh.text_frame.text.strip()[:14]))
        if sh.has_table:
            need = table_need(sh)
            bottom = T + need
            if bottom > 6.92:
                issues.append(f"P{i} 表格過高：底部到 {bottom:.2f}\"（上限 6.92）")
    for a in range(len(boxes)):
        for b in range(a + 1, len(boxes)):
            x1, y1, w1, h1, t1 = boxes[a]; x2, y2, w2, h2, t2 = boxes[b]
            ox = min(x1 + w1, x2 + w2) - max(x1, x2)
            oy = min(y1 + h1, y2 + h2) - max(y1, y2)
            if ox > 0.12 and oy > 0.12:
                issues.append(f"P{i} 文字框重疊 {ox:.2f}x{oy:.2f}\"：「{t1}」/「{t2}」")

print(f"總頁數：{len(prs.slides.__iter__.__self__._sldIdLst)}")
print(f"問題數：{len(issues)}")
for x in issues: print(" -", x)
